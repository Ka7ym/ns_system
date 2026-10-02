import argparse
import re
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

from config import TEMPLATES_DIR


OUTPUTS = {
    "employment": ("employment_contract_sample_v2.docx", "Трудовой договор"),
    "liability": ("material_responsibility_sample_v1.docx", "Договор материальной ответственности"),
    "order": ("employment_order_sample_v2.docx", "Приказ о приёме на работу"),
}


def iter_container(container):
    for paragraph in getattr(container, "paragraphs", []):
        yield paragraph
    for table in getattr(container, "tables", []):
        for row in table.rows:
            for cell in row.cells:
                yield from iter_container(cell)


def all_paragraphs(document):
    yield from iter_container(document)
    for section in document.sections:
        yield from iter_container(section.header)
        yield from iter_container(section.footer)


def set_paragraph_text(paragraph: Paragraph, text: str):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def replace_everywhere(document, old: str, new: str):
    if not old:
        return
    for paragraph in all_paragraphs(document):
        current = paragraph.text
        updated = current.replace(old, new)
        if updated != current:
            set_paragraph_text(paragraph, updated)


def get_employee_name(document, kind: str) -> str:
    paragraphs = [paragraph.text for paragraph in all_paragraphs(document)]
    if kind == "order":
        for text in paragraphs:
            match = re.search(r"Принять\s+(.+?)\s+на работу", text, re.IGNORECASE)
            if match:
                return match.group(1).strip(" .,:;\n")
    else:
        for text in paragraphs:
            match = re.search(
                r"Республики\s+Казахстан\s+(.+?)(?:\s*\(-ая\)|,\s*именуем)",
                text,
                re.IGNORECASE,
            )
            if match:
                return match.group(1).strip(" .,:;\n")
    raise ValueError(f"Не удалось определить ФИО в образце типа {kind}")


def get_employee_iin(document):
    for paragraph in all_paragraphs(document):
        match = re.search(r"(?:ИИН|ЖСН)\s*[:№]?\s*(\d{12})", paragraph.text, re.IGNORECASE)
        if match:
            return match.group(1)
    return None


def replace_order(document, employee_name: str):
    surname = employee_name.split()[0]
    for paragraph in all_paragraphs(document):
        text = paragraph.text
        if re.search(r"^\s*ПРИКАЗ\s*№", text, re.IGNORECASE):
            set_paragraph_text(paragraph, "ПРИКАЗ № {{ORDER_NUMBER}}")
        elif re.search(r"^\s*от\s+\d{2}\.\d{2}\.\d{4}", text, re.IGNORECASE):
            set_paragraph_text(paragraph, "от {{TODAY_DMY}} г.")
        elif "Принять" in text:
            set_paragraph_text(
                paragraph,
                "Принять {{FULL_NAME}}, ИИН {{IIN}}, на должность {{POSITION}} с {{START_DATE_RU}} г. с окладом {{SALARY}} тенге.",
            )
        elif "ОСНОВАНИЕ:" in text:
            set_paragraph_text(paragraph, "ОСНОВАНИЕ: заявление {{FULL_NAME}}.")
        elif "С приказом ознакомлен" in text:
            set_paragraph_text(paragraph, "С приказом ознакомлен(-а): _______________ {{FULL_NAME}}")
        else:
            replace_everywhere(document, employee_name, "{{FULL_NAME}}")
            if surname != employee_name:
                replace_everywhere(document, surname, "{{FULL_NAME}}")


def prepare_template(source: Path, kind: str, output: Path):
    document = Document(source)
    employee_name = get_employee_name(document, kind)
    employee_iin = get_employee_iin(document)

    if kind == "order":
        replace_order(document, employee_name)
    else:
        replace_everywhere(document, employee_name, "{{FULL_NAME}}")
        surname = employee_name.split()[0]
        if surname != employee_name:
            replace_everywhere(document, surname, "{{FULL_NAME}}")
        if employee_iin:
            replace_everywhere(document, employee_iin, "{{IIN}}")

        texts = [paragraph.text for paragraph in all_paragraphs(document)]
        identity_values = set()
        birth_dates = set()
        for text in texts:
            identity = re.search(
                r"Удостоверение личности:\s*№\s*([\d ]+)\s*,?\s*выдано\s*(\d{2}\.\d{2}\.\d{4})",
                text,
                re.IGNORECASE,
            )
            if identity:
                identity_values.add(identity.group(1).strip())
                identity_values.add(identity.group(2))
            if re.search(r"Дата рождения", text, re.IGNORECASE):
                birth_dates.update(re.findall(r"\d{2}\.\d{2}\.\d{4}", text))

        for value in identity_values:
            replace_everywhere(document, value, "{{DOCUMENT_NUMBER}}" if len(re.sub(r"\D", "", value)) >= 8 and "." not in value else "{{DOCUMENT_ISSUE_DATE_RU}}")
        for value in birth_dates:
            replace_everywhere(document, value, "{{BIRTH_DATE}}")

        for paragraph in list(all_paragraphs(document)):
            text = paragraph.text
            if kind == "employment" and re.search(r"^\s*ТРУДОВОЙ ДОГОВОР\s*№", text, re.IGNORECASE):
                set_paragraph_text(paragraph, re.sub(r"№\s*\S+", "№ {{CONTRACT_NUMBER}}", text, count=1))
            if kind == "employment" and "Настоящий Договор заключен сроком на один год" in text:
                updated = re.sub(
                    r"Настоящий Договор заключен сроком на один год с .+? \(включительно\)",
                    "Настоящий Договор {{CONTRACT_TERM_CLAUSE}}",
                    text,
                    count=1,
                )
                set_paragraph_text(paragraph, updated)
            elif kind == "employment" and "Осы Шарт бір жыл мерзімге" in text:
                updated = re.sub(
                    r"Осы Шарт бір жыл мерзімге .+? \(қоса алғанда\) жасалды",
                    "Осы Шарт {{CONTRACT_TERM_CLAUSE_KZ}}",
                    text,
                    count=1,
                )
                set_paragraph_text(paragraph, updated)
            elif kind == "employment" and "Размер месячной заработной платы работника устанавливается" in text:
                set_paragraph_text(paragraph, text + " Оклад по настоящему Договору: {{SALARY}} тенге в месяц.")
            elif kind == "employment" and "Қызметкердің айлық жалақысының мөлшері" in text:
                set_paragraph_text(paragraph, text + " Осы шарт бойынша айлық оклад: {{SALARY}} теңге.")
            elif kind == "employment" and "хат-хабар маманы" in text:
                updated = text.replace(
                    "лауазымындағы жұмысты (еңбек функциясын) жеке өзі орындауға, хат-хабар маманы {{POSITION}} сақтауға",
                    "{{POSITION}} лауазымындағы жұмысты (еңбек функциясын) жеке өзі орындауға",
                )
                set_paragraph_text(paragraph, updated)
            elif re.search(r"^\s*г\.\s*Караганда", text) and re.search(r"\d{1,2}.*\d{4}", text):
                updated = re.sub(r"«\s*\d{1,2}\s*»\s+[А-Яа-яЁё]+\s+\d{4}\s*(?:г\.|ж\.)?", "{{TODAY_RU}}", text, count=1)
                if updated != text:
                    set_paragraph_text(paragraph, updated)
            elif re.search(r"^\s*Қарағанды\s+қ\.", text, re.IGNORECASE):
                updated = re.sub(r"«\s*\d{1,2}\s*»\s+[А-Яа-яЁёӘәҒғҚқҢңӨөҰұҮүҺһІі]+\s+\d{4}\s*(?:г\.|ж\.)?", "{{TODAY_KZ}}", text, count=1)
                if updated != text:
                    set_paragraph_text(paragraph, updated)
            if "Адрес фактического проживания:" in text:
                set_paragraph_text(paragraph, text.replace("Адрес фактического проживания:", "Адрес фактического проживания: {{ADDRESS}}"))
            elif "Адрес регистрации:" in text:
                set_paragraph_text(paragraph, text.replace("Адрес регистрации:", "Адрес регистрации: {{ADDRESS}}"))
            elif re.fullmatch(r"\s*Тел\.:\s*", text):
                set_paragraph_text(paragraph, "Тел.: {{PHONE}}")

        if kind == "employment":
            for paragraph in list(all_paragraphs(document)):
                text = paragraph.text
                position = re.search(r"в должности\s+(.+?)\s+соблюдать", text, re.IGNORECASE)
                if position:
                    set_paragraph_text(paragraph, text.replace(position.group(1), "{{POSITION}}", 1))
                if "хат-хабар маманы" in text.casefold():
                    updated = re.sub(
                        r"^Шарт бойынша Қызметкер.+?сақтауға,",
                        "Шарт бойынша Қызметкер сыйақы (жалақы) үшін {{POSITION}} лауазымындағы жұмысты (еңбек функциясын) жеке өзі орындауға, еңбек тәртібін сақтауға,",
                        text,
                        count=1,
                        flags=re.IGNORECASE,
                    )
                    set_paragraph_text(paragraph, updated)
        else:
            for paragraph in list(all_paragraphs(document)):
                text = paragraph.text
                if "занимающий на основании трудового договора" in text:
                    updated = re.sub(r"занимающий на основании трудового договора\s+.+?\s+в целях", "занимающий должность {{POSITION}} в целях", text, flags=re.IGNORECASE)
                    set_paragraph_text(paragraph, updated)
                elif "атқаратын" in text and "лауазым" in text.casefold():
                    updated = re.sub(r"қосалқы жұмысшы", "{{POSITION}}", text, flags=re.IGNORECASE)
                    set_paragraph_text(paragraph, updated)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)

    verification = "\n".join(paragraph.text for paragraph in all_paragraphs(document))
    if employee_name in verification or (employee_iin and employee_iin in verification):
        raise ValueError(f"Персональные данные исходного образца остались в шаблоне: {source.name}")


def main():
    parser = argparse.ArgumentParser(description="Подготовить обезличенные DOCX-шаблоны найма NS SYSTEM")
    parser.add_argument("--source-dir", type=Path, default=Path.home() / "Desktop")
    args = parser.parse_args()
    candidates = list(args.source_dir.glob("*.docx"))
    sources = {}
    for path in candidates:
        try:
            document = Document(path)
        except Exception:
            continue
        text = "\n".join(paragraph.text for paragraph in all_paragraphs(document))
        lowered = text.casefold()
        if "полной материальной ответственности" in lowered:
            sources["liability"] = path
        elif "приказываю:" in lowered and "на работу" in lowered:
            sources["order"] = path
        elif "трудовой договор" in lowered and "испытательный срок" in lowered:
            sources["employment"] = path

    missing = set(OUTPUTS) - sources.keys()
    if missing:
        raise SystemExit(f"Не найдены исходные документы типов: {', '.join(sorted(missing))}")

    for kind, source in sources.items():
        filename, _ = OUTPUTS[kind]
        destination = Path(TEMPLATES_DIR) / filename
        prepare_template(source, kind, destination)
        print(f"Подготовлен шаблон: {filename}")


if __name__ == "__main__":
    main()
