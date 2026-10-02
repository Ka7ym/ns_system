import datetime
import calendar
import os
from typing import Iterable, Optional

from docx import Document as DocxDocument
from fastapi import HTTPException
from sqlalchemy.orm import Session

from config import DOCS_DIR, TEMPLATES_DIR, load_app_settings
from models import (
    DOCUMENT_STATUS_READY,
    TEMPLATE_APPLICATION,
    TEMPLATE_MATERIAL_RESPONSIBILITY,
    TEMPLATE_TERMINATION,
    Document,
    DocumentTemplate,
    Employee,
)

PLACEHOLDERS = (
    "FULL_NAME",
    "IIN",
    "BIRTH_DATE",
    "DOCUMENT_NUMBER",
    "POSITION",
    "DEPARTMENT",
    "SALARY",
    "START_DATE",
    "START_DATE_RU",
    "CONTRACT_END_DATE_RU",
    "CONTRACT_TERM_CLAUSE",
    "CONTRACT_TERM_CLAUSE_KZ",
    "TODAY_RU",
    "TODAY_KZ",
    "TODAY_DMY",
    "CONTRACT_NUMBER",
    "ORDER_NUMBER",
    "CONTRACT_TYPE",
    "WORK_SCHEDULE",
    "TERMINATION_DATE",
    "TERMINATION_REASON",
    "PHONE",
    "EMAIL",
    "ADDRESS",
    "DOCUMENT_ISSUE_DATE_RU",
    "COMPANY_NAME",
    "COMPANY_BIN",
    "COMPANY_ADDRESS",
    "COMPANY_DIRECTOR",
    "TODAY",
)

TYPE_LABELS = {
    "contract": "Трудовой договор",
    "order": "Приказ о приёме",
    TEMPLATE_MATERIAL_RESPONSIBILITY: "Договор материальной ответственности",
    "consent": "Согласие на обработку персональных данных",
    "application": "Заявление о приёме на работу",
    "other": "Другой документ",
    "termination_application": "Заявление на увольнение",
}


def employee_context(employee: Employee, extra: Optional[dict] = None) -> dict:
    settings = load_app_settings()
    today = datetime.date.today()
    start_date = _parse_date(employee.start_date)
    contract_end = ""
    contract_term = "заключен на неопределенный срок"
    contract_term_kz = "белгісіз мерзімге жасалды"
    if start_date:
        year = start_date.year + 1
        contract_end_date = datetime.date(year, start_date.month, min(start_date.day, calendar.monthrange(year, start_date.month)[1])) - datetime.timedelta(days=1)
        contract_end = _date_ru(contract_end_date)
        start_date_kz = _date_kz(start_date)
        if "сроч" in (employee.contract_type or "").casefold():
            contract_term = f"заключен сроком на один год с {_date_ru(start_date)} по {contract_end} (включительно)"
            contract_term_kz = f"бір жыл мерзімге {start_date_kz} бастап {_date_kz(contract_end_date)} дейін (қоса алғанда) жасалды"
        else:
            contract_term = f"заключен на неопределенный срок с {_date_ru(start_date)}"
            contract_term_kz = f"белгісіз мерзімге {start_date_kz} бастап жасалды"
    contract_number = f"NS-{employee.id}/{today.year}"
    order_number = f"{employee.id}-л/с"
    birth = (
        getattr(employee, "birth_date_full", None)
        or getattr(employee, "birth_full", None)
        or (str(employee.birth_year) if employee.birth_year else "")
    )
    payload = {
        "FULL_NAME": employee.full_name or "",
        "IIN": employee.iin or "",
        "BIRTH_DATE": birth,
        "DOCUMENT_NUMBER": employee.document_number or "",
        "POSITION": employee.position or "",
        "DEPARTMENT": employee.department or "",
        "SALARY": f"{employee.salary:,.0f}".replace(",", " ") if employee.salary is not None else "",
        "START_DATE": employee.start_date or "",
        "START_DATE_RU": _date_ru(start_date) if start_date else employee.start_date or "",
        "CONTRACT_END_DATE_RU": contract_end,
        "CONTRACT_TERM_CLAUSE": contract_term,
        "CONTRACT_TERM_CLAUSE_KZ": contract_term_kz,
        "TODAY_RU": f"«{today.day:02d}» {_RU_MONTHS[today.month - 1]} {today.year} г.",
        "TODAY_KZ": f"«{today.day:02d}» {_KZ_MONTHS[today.month - 1]} {today.year} ж.",
        "TODAY_DMY": today.strftime("%d.%m.%Y"),
        "CONTRACT_NUMBER": contract_number,
        "ORDER_NUMBER": order_number,
        "CONTRACT_TYPE": employee.contract_type or "",
        "WORK_SCHEDULE": employee.work_schedule or "",
        "TERMINATION_DATE": employee.termination_date or "",
        "TERMINATION_REASON": employee.termination_reason or "",
        "PHONE": employee.phone or "",
        "EMAIL": employee.email or "",
        "ADDRESS": employee.address or "",
        "DOCUMENT_ISSUE_DATE_RU": _date_ru(_parse_date(employee.document_issue_date)) or employee.document_issue_date or "",
        "COMPANY_NAME": settings.get("company_name") or "ТОО «НС Система»",
        "COMPANY_BIN": settings.get("company_bin") or "",
        "COMPANY_ADDRESS": settings.get("company_address") or "",
        "COMPANY_DIRECTOR": settings.get("company_director") or "",
        "TODAY": datetime.datetime.now().strftime("%d.%m.%Y"),
    }
    if extra:
        payload.update(extra)
    return payload


_RU_MONTHS = (
    "января", "февраля", "марта", "апреля", "мая", "июня",
    "июля", "августа", "сентября", "октября", "ноября", "декабря",
)


def _parse_date(value: Optional[str]) -> Optional[datetime.date]:
    if not value:
        return None
    value = str(value).strip()
    for date_format in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            return datetime.datetime.strptime(value[:10], date_format).date()
        except ValueError:
            continue
    return None


def _date_ru(value: Optional[datetime.date]) -> str:
    if not value:
        return ""
    return f"{value.day:02d} { _RU_MONTHS[value.month - 1]} {value.year}"


_KZ_MONTHS = (
    "қаңтар", "ақпан", "наурыз", "сәуір", "мамыр", "маусым",
    "шілде", "тамыз", "қыркүйек", "қазан", "қараша", "желтоқсан",
)


def _date_kz(value: datetime.date) -> str:
    return f"{value.year} жылғы {value.day:02d} {_KZ_MONTHS[value.month - 1]}"


def replace_placeholders(doc: DocxDocument, values: dict) -> None:
    def replace_in_paragraph(paragraph) -> None:
        full = "".join(run.text for run in paragraph.runs) if paragraph.runs else paragraph.text
        updated = full
        for key, value in values.items():
            updated = updated.replace("{{" + key + "}}", str(value))
        if updated != full:
            if paragraph.runs:
                paragraph.runs[0].text = updated
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.text = updated

    def replace_in_container(container) -> None:
        for paragraph in getattr(container, "paragraphs", []):
            replace_in_paragraph(paragraph)
        for table in getattr(container, "tables", []):
            for row in table.rows:
                for cell in row.cells:
                    replace_in_container(cell)

    replace_in_container(doc)
    for section in doc.sections:
        replace_in_container(section.header)
        replace_in_container(section.footer)


def create_placeholder_docx(path: str, title: str, extra_lines: Iterable[str]) -> None:
    doc = DocxDocument()
    doc.add_heading(title, 0)
    doc.add_paragraph("ТОО «НС Система»")
    doc.add_paragraph("Дата: {{TODAY}}")
    doc.add_paragraph("ФИО: {{FULL_NAME}}")
    doc.add_paragraph("ИИН: {{IIN}}")
    for line in extra_lines:
        doc.add_paragraph(line)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    doc.save(path)


def ensure_builtin_templates(db: Session, created_by: Optional[int] = None) -> None:
    builtins = [
        (
            "contract",
            "Трудовой договор (базовый)",
            "v1",
            [
                "Должность: {{POSITION}}",
                "Отдел: {{DEPARTMENT}}",
                "Оклад: {{SALARY}} ₸",
                "Дата начала работы: {{START_DATE}}",
                "Тип договора: {{CONTRACT_TYPE}}",
                "График: {{WORK_SCHEDULE}}",
            ],
        ),
        (
            "order",
            "Приказ о приёме (базовый)",
            "v1",
            [
                "Принять {{FULL_NAME}} на должность {{POSITION}}",
                "Подразделение: {{DEPARTMENT}}",
                "Дата приёма: {{START_DATE}}",
                "Оклад: {{SALARY}} ₸",
            ],
        ),
        (
            "termination_application",
            "Заявление на увольнение (базовый)",
            "v1",
            [
                "Прошу уволить меня с должности {{POSITION}} ({{DEPARTMENT}})",
                "Дата увольнения: {{TERMINATION_DATE}}",
                "Причина: {{TERMINATION_REASON}}",
            ],
        ),
        (
            "consent",
            "Согласие на обработку ПДн (базовый)",
            "v1",
            [
                "Я, {{FULL_NAME}}, ИИН {{IIN}}, даю согласие на обработку персональных данных.",
            ],
        ),
        (
            TEMPLATE_APPLICATION,
            "Заявление о приёме на работу (базовый)",
            "v1",
            [
                "Прошу принять меня на работу в ТОО «НС Система».",
                "ФИО: {{FULL_NAME}}",
                "Должность: {{POSITION}}",
                "Дата приёма: {{START_DATE}}",
            ],
        ),
    ]
    for doc_type, name, version, lines in builtins:
        exists = (
            db.query(DocumentTemplate)
            .filter(DocumentTemplate.type == doc_type, DocumentTemplate.version == version)
            .first()
        )
        if exists:
            continue
        filename = f"{doc_type}_{version}.docx"
        path = os.path.join(TEMPLATES_DIR, filename)
        create_placeholder_docx(path, name, lines)
        has_active = (
            db.query(DocumentTemplate)
            .filter(DocumentTemplate.type == doc_type, DocumentTemplate.is_active == True)  # noqa: E712
            .first()
        )
        db.add(
            DocumentTemplate(
                name=name,
                type=doc_type,
                description="Встроенный шаблон NS SYSTEM",
                file_path=path,
                version=version,
                is_active=has_active is None,
                created_by=created_by,
            )
        )
    imported_templates = [
        ("contract", "Трудовой договор (шаблон компании)", "v2", "employment_contract_sample_v2.docx"),
        ("order", "Приказ о приёме (шаблон компании)", "v2", "employment_order_sample_v2.docx"),
        (TEMPLATE_MATERIAL_RESPONSIBILITY, "Договор материальной ответственности", "v1", "material_responsibility_sample_v1.docx"),
    ]
    for doc_type, name, version, filename in imported_templates:
        path = os.path.join(TEMPLATES_DIR, filename)
        if not os.path.exists(path):
            continue
        exists = db.query(DocumentTemplate).filter(
            DocumentTemplate.type == doc_type,
            DocumentTemplate.version == version,
        ).first()
        if exists:
            continue
        db.query(DocumentTemplate).filter(DocumentTemplate.type == doc_type).update({"is_active": False})
        db.add(
            DocumentTemplate(
                name=name,
                type=doc_type,
                description="Обезличенный шаблон, импортированный из предоставленного документа",
                file_path=path,
                version=version,
                is_active=True,
                created_by=created_by,
            )
        )
    db.commit()


class DocumentService:
    @staticmethod
    def get_active_template(db: Session, doc_type: str) -> Optional[DocumentTemplate]:
        return (
            db.query(DocumentTemplate)
            .filter(
                DocumentTemplate.type == doc_type,
                DocumentTemplate.is_active == True,  # noqa: E712
            )
            .order_by(DocumentTemplate.id.desc())
            .first()
        )

    @classmethod
    def generate_selected(
        cls,
        db: Session,
        employee: Employee,
        types: list[str],
        created_by: Optional[int] = None,
        extra: Optional[dict] = None,
        require_template: bool = False,
    ) -> list[Document]:
        created: list[Document] = []
        for doc_type in types:
            template = cls.get_active_template(db, doc_type)
            if not template:
                if require_template or doc_type == "termination_application":
                    raise HTTPException(
                        status_code=409,
                        detail=f"Шаблон «{TYPE_LABELS.get(doc_type, doc_type)}» не загружен. Загрузите шаблон в разделе «Шаблоны документов».",
                    )
                continue
            if not os.path.exists(template.file_path):
                raise HTTPException(status_code=409, detail="Файл шаблона не найден на сервере")
            timestamp = datetime.datetime.now().strftime("%Y%m%d%H%M%S")
            filename = f"{doc_type}_{employee.id}_{timestamp}.docx"
            output_path = os.path.join(DOCS_DIR, filename)
            doc = DocxDocument(template.file_path)
            replace_placeholders(doc, employee_context(employee, extra))
            doc.save(output_path)
            record = Document(
                employee_id=employee.id,
                template_id=template.id,
                type=doc_type,
                file_name=filename,
                file_path=output_path,
                status=DOCUMENT_STATUS_READY,
                version=template.version,
                created_by=created_by,
            )
            db.add(record)
            created.append(record)
        if not created:
            raise HTTPException(
                status_code=409,
                detail="Нет активных шаблонов для выбранных документов. Загрузите шаблоны в разделе «Шаблоны документов».",
            )
        db.flush()
        return created

    @classmethod
    def generate_documents(
        cls,
        db: Session,
        employee: Employee,
        types: Optional[list[str]] = None,
        created_by: Optional[int] = None,
        user_id: Optional[int] = None,
        require_templates: bool = False,
        extra: Optional[dict] = None,
    ) -> list[Document]:
        return cls.generate_selected(
            db,
            employee,
            types or ["contract", "order"],
            created_by=created_by or user_id,
            extra=extra,
            require_template=require_templates,
        )
